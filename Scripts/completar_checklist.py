"""
Completa automáticamente el checklist de un expediente.

Tú rellenas en el checklist (columnas con título AMARILLO en la plantilla):
    Nº justificante, Proveedor, Nº factura, Importe (€), Forma de pago,
    Fecha pago  (y, si quieres, Archivo original).

Este script rellena lo demás (columnas con título GRIS), SOLO en las celdas
que estén vacías, así que no borra nada de lo que hayas escrito tú:
    - Archivo original: busca el PDF de cada factura por su nº de factura
      (en el nombre del archivo o dentro del PDF).
    - Tipo: DIGITAL o ESCANEADA (papel), analizando cada PDF.
    - TEXTO EXACTO DEL SELLO: a partir de "texto_sello" del config.json.
    - Copia sellada: nombre del PDF sellado, empezando por el nº de justificante.
    - Qué hacer.
    - Pago localizado: busca cada pago en el extracto bancario (mismo importe,
      fecha cercana; también cobros conjuntos de varias facturas y diferencias
      de céntimos).

    python Scripts/completar_checklist.py "NOMBRE DEL EXPEDIENTE"
    python Scripts/completar_checklist.py "NOMBRE DEL EXPEDIENTE" --rehacer
        (--rehacer vuelve a calcular las columnas automáticas aunque ya
         tengan algo; útil si has cambiado importes o fechas)

Guarda una copia de seguridad del checklist antes de modificarlo
(en la misma carpeta, terminada en "(copia antes de completar).xlsx").
Cierra el Excel antes de ejecutarlo.
"""
import datetime as dt
import os
import re
import shutil
import sys

import openpyxl
import pymupdf
from openpyxl.styles import PatternFill

import banco
import expediente

AMARILLO = PatternFill("solid", fgColor="FFE699")
ROJO = PatternFill("solid", fgColor="F8CBAD")
SIN_RELLENO = PatternFill(fill_type=None)
VENTANA_DIAS = 10  # margen entre la fecha de pago anotada y la del banco

TEXTO_POR_DEFECTO = "ESTA FACTURA ESTA SUBVENCIONADA POR ...; IMPORTE: {importe} € ({porcentaje})"


def norm(s):
    return re.sub(r"[^0-9a-z]", "", str(s).lower())


def tokens(texto):
    return {norm(t) for t in re.split(r"\s+", str(texto)) if norm(t)} | \
           {norm(t) for t in re.split(r"[^0-9A-Za-z]+", str(texto)) if norm(t)}


def a_fechas(v):
    """Celda de fecha(s) -> lista de date. Admite fecha de Excel o texto
    'dd/mm/aaaa', y varias separadas por ';' o ' / '."""
    if v is None or v == "":
        return []
    if isinstance(v, dt.datetime):
        return [v.date()]
    if isinstance(v, dt.date):
        return [v]
    res = []
    for parte in re.split(r";|\s/\s", str(v)):
        m = re.search(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", parte)
        if m:
            d, mes, a = (int(x) for x in m.groups())
            res.append(dt.date(a + 2000 if a < 100 else a, mes, d))
    return res


# ---------------------------------------------------------------- facturas

def analizar_pdf(ruta):
    """(tipo, texto) de un PDF. ESCANEADA si alguna página es básicamente una
    imagen a página completa sin texto real (o con OCR encima)."""
    doc = pymupdf.open(ruta)
    texto = " ".join(p.get_text() for p in doc)
    escaneada = False
    for p in doc:
        area = p.rect.width * p.rect.height
        cobertura = 0
        for img in p.get_images(full=True):
            for r in p.get_image_rects(img[0]):
                cobertura = max(cobertura, (r & p.rect).get_area() / area)
        productor = (doc.metadata.get("producer", "") + doc.metadata.get("creator", "")).lower()
        if cobertura > 0.85 and (len(p.get_text().strip()) < 50 or "scan" in productor):
            escaneada = True
    return ("ESCANEADA (papel)" if escaneada else "DIGITAL"), texto


def indexar_facturas(carpeta):
    """{ruta relativa: {'tipo', 'tokens', 'tokens_nombre'}} de todos los PDF."""
    idx = {}
    for raiz, _, archivos in os.walk(carpeta):
        for a in archivos:
            if a.lower().endswith(".pdf"):
                ruta = os.path.join(raiz, a)
                rel = os.path.relpath(ruta, carpeta)
                try:
                    tipo, texto = analizar_pdf(ruta)
                except Exception as e:  # PDF dañado o protegido
                    print(f"  (no se puede leer {rel}: {e})")
                    continue
                idx[rel] = {"tipo": tipo, "tokens": tokens(texto), "tokens_nombre": tokens(os.path.splitext(a)[0])}
    return idx


NO_FACTURA = ("justificante", "albaran", "albarán", "pedido", "presupuesto", "ticket")


def buscar_archivo(nfactura, idx):
    """Devuelve (ruta relativa o None, lista de otras coincidencias)."""
    n = norm(nfactura)
    if len(n) < 3:
        return None, []

    def coincide(toks):
        return n in toks or (len(n) >= 6 and any(n in t for t in toks))

    por_nombre = [r for r, d in idx.items() if coincide(d["tokens_nombre"])]
    por_texto = [r for r, d in idx.items() if r not in por_nombre and coincide(d["tokens"])]
    candidatos = por_nombre + por_texto
    # mejor las que no parecen justificantes, albaranes, pedidos...
    candidatos.sort(key=lambda r: any(p in r.lower() for p in NO_FACTURA))
    if not candidatos:
        return None, []
    return candidatos[0], candidatos[1:]


# ---------------------------------------------------------------- pagos

def buscar_pagos(pendientes, movimientos, usados):
    """pendientes: lista de dicts con num, proveedor, importe, fecha.
    Devuelve {num: descripción}. Tres pasadas: pago exacto, cobro conjunto
    de varias facturas del mismo proveedor y fecha, y diferencia de céntimos."""
    resultado = {}

    def candidato(importe, fecha, tolerancia=0.005):
        opciones = []
        for k, m in enumerate(movimientos):
            if k in usados or m["cargo"] is False or abs(m["importe"] - importe) >= tolerancia:
                continue
            dias = abs((m["fecha"] - fecha).days) if fecha else 0
            if fecha and dias > VENTANA_DIAS:
                continue
            opciones.append((m["cargo"] is not True, dias, k))  # mejor con signo "-" y fecha cercana
        if not opciones or (not fecha and len(opciones) > 1):
            return None
        k = min(opciones)[2]
        usados.add(k)
        m = movimientos[k]
        txt = banco.descripcion(m)
        if abs(m["importe"] - importe) >= 0.005:
            txt += f" — OJO: difiere {banco.euros(abs(m['importe'] - importe))} € del importe"
        return txt

    for p in pendientes:
        resultado[p["num"]] = candidato(p["importe"], p["fecha"])
    grupos = {}
    for p in pendientes:
        if not resultado[p["num"]] and p["fecha"]:
            grupos.setdefault((norm(p["proveedor"]), p["fecha"]), []).append(p)
    for grupo in grupos.values():
        if len(grupo) > 1:
            txt = candidato(sum(p["importe"] for p in grupo), grupo[0]["fecha"])
            if txt:
                for p in grupo:
                    resultado[p["num"]] = txt + " — cobro conjunto de nº " + " + ".join(q["num"] for q in grupo)
    for p in pendientes:
        if not resultado[p["num"]]:
            resultado[p["num"]] = candidato(p["importe"], p["fecha"], tolerancia=0.02)
    return resultado


# ---------------------------------------------------------------- principal

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    rehacer = "--rehacer" in sys.argv
    if len(args) != 1:
        sys.exit(__doc__)
    cfg = expediente.cargar(args[0])
    ruta = cfg["checklist"]
    try:
        wb = openpyxl.load_workbook(ruta)
    except PermissionError:
        sys.exit("No puedo abrir el checklist: ciérralo en Excel y vuelve a intentarlo.")
    ws = wb.worksheets[0]
    c = expediente.columnas([cell.value for cell in ws[1]])
    faltan = [k for k in ("num", "proveedor", "nfactura", "importe", "texto", "tipo", "archivo", "copia", "pago")
              if c[k] is None]
    if faltan:
        sys.exit(f"Al checklist le faltan columnas: {', '.join(expediente.COL[k] for k in faltan)}.\n"
                 "Usa la plantilla (Expedientes/_PLANTILLA/Checklist.xlsx) y no cambies los títulos.")

    def celda(fila, clave):
        return ws.cell(fila, c[clave] + 1) if c[clave] is not None else None

    def vacia(fila, clave):
        cl = celda(fila, clave)
        return cl is not None and (rehacer or cl.value in (None, ""))

    def nota(fila, texto):
        cl = celda(fila, "notas")
        if cl is not None and texto not in str(cl.value or ""):
            cl.value = (str(cl.value) + "\n" if cl.value else "") + texto

    print("Analizando las facturas de", cfg["facturas_originales"], "...")
    idx = indexar_facturas(cfg["facturas_originales"])
    print(f"  {len(idx)} PDF encontrados.")
    plantilla = cfg.get("texto_sello") or TEXTO_POR_DEFECTO

    filas = [f for f in range(2, ws.max_row + 1) if celda(f, "num").value not in (None, "")]
    usados_archivos, problemas = set(), []
    pendientes_pago, lineas_de_fila = [], {}
    for f in filas:
        num = str(celda(f, "num").value).strip()
        importe = expediente.a_numero(celda(f, "importe").value)
        if importe is None:
            problemas.append(f"Nº {num}: falta el importe.")
            continue

        # archivo original
        if celda(f, "archivo").value in (None, ""):
            archivo, otros = buscar_archivo(celda(f, "nfactura").value, idx)
            if archivo:
                celda(f, "archivo").value = archivo
                nota(f, "[auto] Archivo encontrado automáticamente: comprobar que es la factura correcta.")
                if otros:
                    nota(f, "[auto] Otros archivos con ese nº: " + "; ".join(otros[:3]))
            else:
                problemas.append(f"Nº {num}: no encuentro el PDF de la factura {celda(f, 'nfactura').value}. "
                                 "Escribe su ruta en 'Archivo original'.")
        archivo = celda(f, "archivo").value
        if archivo:
            usados_archivos.add(os.path.normpath(str(archivo)))
            if not os.path.isfile(os.path.join(cfg["facturas_originales"], str(archivo))):
                problemas.append(f"Nº {num}: no existe el archivo '{archivo}'.")

        # tipo, texto, copia, qué hacer
        info = idx.get(os.path.normpath(str(archivo))) if archivo else None
        if info and vacia(f, "tipo"):
            celda(f, "tipo").value = info["tipo"]
        tipo = str(celda(f, "tipo").value or "")
        escaneada = tipo.startswith("ESCANEADA")
        porcentaje = str(celda(f, "porcentaje").value or "100%") if c.get("porcentaje") is not None else "100%"
        if vacia(f, "texto"):
            celda(f, "texto").value = plantilla.format(importe=banco.euros(importe), porcentaje=porcentaje)
        nums = [n.strip() for n in num.split(",")]
        prefijo = f"{int(nums[0]):02d}" if nums[0].isdigit() else nums[0]
        prefijo += "".join("-" + n for n in nums[1:])
        if vacia(f, "copia") and tipo:
            if escaneada:
                nombre = f"{prefijo} - {celda(f, 'proveedor').value} {celda(f, 'nfactura').value}.pdf"
            else:
                nombre = f"{prefijo} - {os.path.basename(str(archivo))}"
            celda(f, "copia").value = re.sub(r'[\\/:*?"<>|]', "-", nombre)
        if c["que_hacer"] is not None and vacia(f, "que_hacer") and tipo:
            celda(f, "que_hacer").value = ("Sellar a mano el original y escanear en color" if escaneada
                                           else "Sellar con el programa")
        for cl in ws[f]:
            cl.fill = AMARILLO if escaneada else SIN_RELLENO

        # pagos pendientes de buscar
        if vacia(f, "pago"):
            forma = str(celda(f, "forma_pago").value or "") if c["forma_pago"] is not None else ""
            fechas = a_fechas(celda(f, "fecha_pago").value) if c["fecha_pago"] is not None else []
            importes = []
            if c["importes_pagos"] is not None and celda(f, "importes_pagos").value:
                importes = [expediente.a_numero(x) for x in re.split(r";|\s/\s", str(celda(f, "importes_pagos").value))]
            if not importes:
                importes = [importe]
            if "efectivo" in forma.lower():
                celda(f, "pago").value = "(efectivo: no aparece en el extracto; aportar recibo)"
                continue
            if len(importes) > 1 and len(fechas) < len(importes):
                problemas.append(f"Nº {num}: hay {len(importes)} pagos pero no todas sus fechas.")
            lineas_de_fila[f] = []
            for i, imp in enumerate(importes):
                clave = f"{f}:{i}"
                lineas_de_fila[f].append(clave)
                pendientes_pago.append({"num": clave, "proveedor": celda(f, "proveedor").value, "importe": imp,
                                        "fecha": fechas[i] if i < len(fechas) else (fechas[0] if fechas else None)})

    # ---- pagos en el extracto
    if pendientes_pago:
        if not cfg.get("extracto") or not os.path.isfile(cfg["extracto"]):
            problemas.append("No encuentro el extracto bancario indicado en config.json: no se buscan pagos.")
        else:
            print("Buscando los pagos en", cfg["extracto"], "...")
            movimientos, _ = banco.leer_movimientos(cfg["extracto"])
            if not movimientos:
                problemas.append("No reconozco movimientos en el extracto (¿es un PDF escaneado?).")
            else:
                usados = set()
                # los ya anotados en el checklist no se vuelven a usar
                for f in filas:
                    if f not in lineas_de_fila and celda(f, "pago").value:
                        for m in banco.RE_DESCRIPCION.finditer(str(celda(f, "pago").value)):
                            for k, mov in enumerate(movimientos):
                                if (str(mov["pag"]), mov["fecha_txt"], banco.euros(mov["importe"])) == \
                                        (m.group(3), m.group(1), m.group(2)):
                                    usados.add(k)
                encontrados = buscar_pagos(pendientes_pago, movimientos, usados)
                # traducir claves internas "fila:i" a nº de justificante en el texto
                nums_de = {f"{f}:{i}": str(celda(f, "num").value).split(", ")[0]
                           for f, ks in lineas_de_fila.items() for i, _ in enumerate(ks)}
                for f, claves in lineas_de_fila.items():
                    textos = []
                    for k in claves:
                        t = encontrados.get(k) or "NO LOCALIZADO: buscar a mano"
                        t = re.sub(r"cobro conjunto de nº (.*)$",
                                   lambda m: "cobro conjunto de nº " + " + ".join(nums_de.get(x, x) for x in m.group(1).split(" + ")), t)
                        textos.append(t)
                    celda(f, "pago").value = "\n".join(textos)
                    celda(f, "pago").fill = ROJO if any(t.startswith("NO LOC") for t in textos) else \
                        (AMARILLO if str(celda(f, "tipo").value or "").startswith("ESCANEADA") else SIN_RELLENO)

    # ---- guardar (con copia de seguridad)
    copia = ruta[:-5] + " (copia antes de completar).xlsx"
    shutil.copy2(ruta, copia)
    try:
        wb.save(ruta)
    except PermissionError:
        sys.exit("No puedo guardar el checklist: ciérralo en Excel y vuelve a intentarlo.")

    # ---- resumen
    total = len(filas)
    tipos = [str(celda(f, "tipo").value or "") for f in filas]
    sin_pago = [str(celda(f, "num").value) for f in filas if str(celda(f, "pago").value or "").startswith("NO LOC")]
    print(f"\nChecklist completado: {total} facturas "
          f"({sum(t.startswith('DIGITAL') for t in tipos)} digitales, "
          f"{sum(t.startswith('ESCANEADA') for t in tipos)} en papel).")
    print(f"Copia de seguridad del checklist anterior: {os.path.basename(copia)}")
    if sin_pago:
        print(f"Pagos NO localizados (en rojo en el Excel): Nº {', '.join(sin_pago)}")
    no_usados = sorted(r for r in idx if os.path.normpath(r) not in usados_archivos)
    if no_usados:
        print(f"\nPDF de la carpeta de facturas que no están en el checklist ({len(no_usados)}); "
              "normalmente albaranes, justificantes o duplicados, pero revisa que no falte ninguna factura:")
        for r in no_usados:
            print("   ", r)
    if problemas:
        print("\nREVISAR:")
        for p in problemas:
            print("  -", p)
    print("\nAbre el checklist y revisa sobre todo las notas '[auto]' y lo marcado en rojo.")


if __name__ == "__main__":
    main()
