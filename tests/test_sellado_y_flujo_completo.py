"""Sellado de un PDF y flujo completo de un expediente (datos inventados)."""
import os
import subprocess
import sys

import numpy as np
import openpyxl
import pymupdf
from PIL import Image

import sellador_facturas as sf
from conftest import RAIZ, pdf_factura


def test_sello_automatico_no_tapa_el_texto_de_la_factura(tmp_path):
    entrada, salida = tmp_path / "f.pdf", tmp_path / "s.pdf"
    pdf_factura(entrada, "123", "10,00", lineas_extra=25)
    originales = [pymupdf.Rect(w[:4]) for w in pymupdf.open(entrada)[0].get_text("words")]
    texto = "ESTA FACTURA ESTA SUBVENCIONADA; IMPORTE: 10,00 € (100%)"
    avisos = sf.sellar_pdf(str(entrada), str(salida), sf.OpcionesSello(texto=texto))
    assert avisos == []
    pagina = pymupdf.open(salida)[0]
    assert texto in " ".join(pagina.get_text().split())
    nuevas = [pymupdf.Rect(w[:4]) for w in pagina.get_text("words")
              if w[4] in ("ESTA", "SUBVENCIONADA;", "IMPORTE:")]
    assert nuevas and not any(n.intersects(o) for n in nuevas for o in originales)


def test_el_original_no_se_modifica(tmp_path):
    entrada = tmp_path / "f.pdf"
    pdf_factura(entrada, "123", "10,00")
    antes = entrada.read_bytes()
    sf.sellar_pdf(str(entrada), str(tmp_path / "s.pdf"), sf.OpcionesSello(texto="SELLO"))
    assert entrada.read_bytes() == antes


def test_fondo_blanco_del_logo_se_vuelve_transparente(tmp_path):
    img = np.full((40, 40, 3), 255, dtype=np.uint8)
    img[10:30, 10:30] = 0  # cuadrado negro en el centro
    ruta = tmp_path / "sello.jpg"
    Image.fromarray(img).save(ruta)
    sf._cargar_logo.cache_clear()
    lector, aspecto = sf._cargar_logo(str(ruta))
    rgba = np.asarray(lector._image)
    assert aspecto == 1.0
    assert rgba[0, 0, 3] == 0        # esquina blanca -> transparente
    assert rgba[20, 20, 3] > 240     # centro negro -> opaco


def ejecutar(script, *args):
    r = subprocess.run([sys.executable, os.path.join(RAIZ, "Scripts", script), *map(str, args)],
                       capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    assert r.returncode == 0, r.stdout + r.stderr
    return r.stdout


def test_flujo_completo_de_un_expediente(expediente_prueba):
    exp = expediente_prueba

    # 1) completar checklist
    ejecutar("completar_checklist.py", exp)
    ws = openpyxl.load_workbook(exp / "Checklist.xlsx").worksheets[0]
    filas = {str(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    # columnas de la plantilla: J archivo, K tipo, L texto, M copia, O pago
    assert filas["1"][9] == "zooplus marzo.pdf"              # no el "justificante pedido..."
    assert filas["1"][10] == "DIGITAL"
    assert filas["1"][11].endswith("IMPORTE: 84,81 € (100%)")
    assert filas["1"][12] == "01 - zooplus marzo.pdf"
    assert "03/03/2025" in filas["1"][14] and "-84,81" in filas["1"][14]
    assert "cobro conjunto de nº 2 + 3" in filas["2"][14]
    assert filas["4"][10].startswith("ESCANEADA")
    assert "-12,50" in filas["4"][14]
    assert filas["5"][14].startswith("(efectivo")
    assert (exp / "Checklist (copia antes de completar).xlsx").exists()

    # 2) sellar: solo las 3 digitales
    salida = ejecutar("sellar_facturas_lote.py", exp)
    assert "3 facturas selladas" in salida
    selladas = sorted(os.listdir(exp / "facturas SELLADAS"))
    assert "01 - zooplus marzo.pdf" in selladas and "_revision" in selladas
    texto = " ".join(pymupdf.open(exp / "facturas SELLADAS" / "01 - zooplus marzo.pdf")[0].get_text().split())
    assert "IMPORTE: 84,81 € (100%)" in texto

    # 3) extracto: solo los pagos, sin datos de terceros
    ejecutar("preparar_extracto_pagos.py", exp)
    limpio = pymupdf.open(exp / "Justificantes de pago" / "Extracto bancario - pagos justificados.pdf")
    t = limpio[0].get_text()
    for visible in ("84,81", "113,00", "12,50", "Nº 1", "Nº 2, 3", "Nº 4"):
        assert visible in t
    for oculto in ("MARIA DONANTE", "SOCIA EJEMPLO", "TIENDA PERSONAL", "40,00"):
        assert oculto not in t

    # 4) entrega: un solo PDF con índice, facturas y justificantes, con marcadores
    ejecutar("preparar_entrega.py", exp)
    entrega = pymupdf.open(exp / "ENTREGA" / "Documentacion justificativa.pdf")
    toc = [t[1] for t in entrega.get_toc()]
    assert toc[0] == "ÍNDICE" and "JUSTIFICANTES DE PAGO" in toc
    assert any(t.startswith("Nº 1 - Zooplus SE") for t in toc)
    assert "84,81 €" in entrega[0].get_text()


def test_completar_avisa_de_errores_en_el_excel_sin_romperse(expediente_prueba):
    """Importe con letras, fecha imposible y un pago que no está en el banco:
    el programa avisa (no se para) y el pago no localizado sigue en rojo al repetir."""
    exp = expediente_prueba
    wb = openpyxl.load_workbook(exp / "Checklist.xlsx")
    ws = wb.worksheets[0]
    for r, fila in enumerate([["6", None, "Proveedor X", "X-1", "abc", None, "Tarjeta", "01/03/2025"],
                              ["7", None, "Proveedor Y", "Y-1", "999,99", None, "Tarjeta", "31/02/2025"],
                              ["8", None, "Proveedor Z", "Z-1", "777,77", None, "Tarjeta", "02/03/2025"]], 7):
        for c, v in enumerate(fila, 1):
            ws.cell(r, c).value = v
    wb.save(exp / "Checklist.xlsx")

    salida = ejecutar("completar_checklist.py", exp)
    assert "Nº 6: el importe 'abc' no se entiende" in salida
    assert "Nº 7: la fecha de pago '31/02/2025' no es válida" in salida
    ejecutar("completar_checklist.py", exp)  # segunda vez: no debe perder el rojo
    ws = openpyxl.load_workbook(exp / "Checklist.xlsx").worksheets[0]
    celda_pago = ws.cell(9, 15)  # fila del nº 8, columna "Pago localizado"
    assert str(celda_pago.value).startswith("NO LOCALIZADO")
    assert celda_pago.fill.fgColor.rgb.endswith("F8CBAD")
